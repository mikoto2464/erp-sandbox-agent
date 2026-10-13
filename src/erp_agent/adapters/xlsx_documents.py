import io
from decimal import Decimal
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from erp_agent.domain.errors import ERPError
from erp_agent.domain.evidence import contextualize_document
from erp_agent.domain.models import Cell, DocumentBlock, Evidence, ParsedDocument, SourceDocument


class XLSXDocumentParser:
    def parse(self, source: SourceDocument, content: bytes) -> ParsedDocument:
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 3000 or sum(e.file_size for e in entries) > 80 * 1024 * 1024:
                    raise ERPError("document_limit", "XLSX 解压内容过大，请拆分文件。")
            workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
        except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError) as exc:
            raise ERPError("invalid_xlsx", "XLSX 无法解析，请检查文件格式或是否加密。") from exc
        blocks: list[DocumentBlock] = []
        evidence: list[Evidence] = []
        warnings: list[str] = []
        cell_count = text_size = 0
        try:
            if len(workbook.worksheets) > 100:
                raise ERPError("document_limit", "XLSX 工作表数量过多，请拆分文件。")
            for sheet in workbook.worksheets:
                cell_count += sheet.max_row * sheet.max_column
                if cell_count > 100000:
                    raise ERPError("document_limit", "XLSX 表格范围过大，请缩小或拆分文件。")
                index = len(blocks)
                eid = f"{source.id}:p{index}"
                title = f"工作表：{sheet.title}" + (
                    "（隐藏）" if sheet.sheet_state != "visible" else ""
                )
                blocks.append(DocumentBlock(index, "paragraph", title, eid))
                evidence.append(Evidence(eid, source.id, f"XLSX 工作表 {sheet.title}", title))
                index = len(blocks)
                anchors = {(r.min_row, r.min_col): r for r in sheet.merged_cells.ranges}
                cells: list[Cell] = []
                formula_count = 0
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell, MergedCell):
                            continue
                        if cell.value is None and not cell.has_style:
                            continue
                        text = "" if cell.value is None else str(cell.value)
                        if cell.data_type == "n" and isinstance(cell.value, (int, float)):
                            text = format(cell.value, ".15g")
                            if "%" in cell.number_format:
                                text = format(Decimal(text) * 100, "f") + "%"
                        text_size += len(text)
                        if text_size > 2000000:
                            raise ERPError("document_limit", "XLSX 提取内容过多，请拆分上传。")
                        formula_count += cell.data_type == "f"
                        merge = anchors.get((cell.row, cell.column))
                        span = merge.max_col - cell.column + 1 if merge else 1
                        end_row = merge.max_row if merge else cell.row
                        for target_row in range(cell.row, end_row + 1):
                            eid = f"{source.id}:t{index}:r{target_row - 1}:c{cell.column - 1}"
                            value = text if target_row == cell.row else ""
                            marker = (
                                ("restart" if end_row > cell.row else "")
                                if target_row == cell.row
                                else "continue"
                            )
                            cells.append(
                                Cell(
                                    target_row - 1,
                                    cell.column - 1,
                                    value,
                                    eid,
                                    span,
                                    marker,
                                    cell.number_format,
                                )
                            )
                            evidence.append(
                                Evidence(
                                    eid,
                                    source.id,
                                    f"XLSX 工作表 {sheet.title} / {get_column_letter(cell.column)}{target_row}",
                                    value,
                                )
                            )
                if formula_count:
                    warnings.append(
                        f"工作表 {sheet.title} 含 {formula_count} 个公式，保留公式原文；读题不重算，公式不作为数值事实。"
                    )
                cells.sort(key=lambda c: (c.row, c.column))
                blocks.append(DocumentBlock(index, "table", cells=tuple(cells)))
        finally:
            workbook.close()
        if not any(c.text.strip() for b in blocks for c in b.cells):
            raise ERPError("empty_document", "XLSX 没有可读取数据。")
        return contextualize_document(
            ParsedDocument(source, tuple(blocks), tuple(evidence), tuple(warnings), "openpyxl-1")
        )
