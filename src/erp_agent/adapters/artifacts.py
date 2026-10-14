import io
import math
import re
import shutil
import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

from erp_agent.adapters.docx_reports import DRAFT_NOTICE, period_label
from erp_agent.adapters.docx_reports import DocxReportRenderer as DocxReportRenderer
from erp_agent.application.ports import FormulaRecalculator, SpreadsheetConverter
from erp_agent.domain.errors import ERPError
from erp_agent.domain.models import BudgetTemplate, RenderedArtifact, ReportData, WorkbookMapping
from erp_agent.domain.workbook import mapped_value, output_sources


class LibreOfficeSpreadsheetConverter:
    def __init__(self, executable: str = 'libreoffice', timeout: int = 60) -> None:
        self.executable, self.timeout = executable, timeout

    def convert(self, content: bytes) -> bytes:
        return self._run(content, 'xls', 'xlsx')

    def _run(self, content: bytes, source_format: str, output_format: str) -> bytes:
        executable = shutil.which(self.executable)
        if not executable:
            raise ERPError('spreadsheet_converter_unavailable', 'XLS 填写需要 LibreOffice；可另存为 XLSX 后上传。', 503)
        with tempfile.TemporaryDirectory(prefix='erp-sheet-') as directory:
            root = Path(directory)
            source = root / ('input.' + source_format)
            source.write_bytes(content)
            output = root / 'output'
            output.mkdir()
            try:
                result = subprocess.run([executable, f'-env:UserInstallation={(root / "profile").as_uri()}',
                                         '--headless', '--convert-to', output_format, '--outdir', str(output), str(source)],
                                        capture_output=True, timeout=self.timeout, check=False)
            except subprocess.TimeoutExpired as exc:
                raise ERPError('spreadsheet_conversion_timeout', '表格转换超时。', 504) from exc
            target = output / ('input.' + output_format)
            if result.returncode or not target.is_file() or target.stat().st_size > 30 * 1024 * 1024:
                raise ERPError('spreadsheet_conversion_failed', '表格转换失败。')
            return target.read_bytes()


class PendingFormulaRecalculator:
    def recalculate(self, content: bytes) -> RenderedArtifact:
        return RenderedArtifact(content, 'pending', ('原公式保持不变，Excel 打开后自动计算。',))


class LibreOfficeFormulaRecalculator:
    def __init__(self, executable: str = 'libreoffice') -> None:
        self.executable = executable

    def recalculate(self, content: bytes) -> RenderedArtifact:
        if not shutil.which(self.executable):
            return PendingFormulaRecalculator().recalculate(content)
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        try:
            if any(re.search(r'\[|WEBSERVICE|HYPERLINK|DDE|RTD|https?://', str(c.value), re.I)
                   for sheet in workbook for row in sheet for c in row if c.data_type == 'f'):
                return RenderedArtifact(content, 'pending', ('模板包含可能引用外部内容的公式，未执行自动重算。',))
        finally:
            workbook.close()
        try:
            result = LibreOfficeSpreadsheetConverter(self.executable)._run(content, 'xlsx', 'xlsx')
        except ERPError:
            return RenderedArtifact(content, 'pending', ('公式自动重算未完成，保留原公式并等待重新计算。',))
        checked = load_workbook(io.BytesIO(result), data_only=True)
        try:
            if any(c.data_type == 'e' for sheet in checked for row in sheet for c in row):
                return RenderedArtifact(content, 'pending', ('公式重算出现错误值，请核对后重新计算。',))
        finally:
            checked.close()
        return RenderedArtifact(result, 'recalculated')


class XlsxWorkbookRenderer:
    def __init__(self, converter: SpreadsheetConverter | None = None,
                 recalculator: FormulaRecalculator | None = None) -> None:
        self.converter = converter
        self.recalculator = recalculator or PendingFormulaRecalculator()

    def render(self, report: ReportData, template: BudgetTemplate | None,
               content: bytes | None, mapping: WorkbookMapping | None) -> RenderedArtifact:
        warnings = [] if report.solution.evaluation.verified else [DRAFT_NOTICE]
        if template:
            if content is None or mapping is None:
                raise ERPError('mapping_required', '请先确认模板映射。')
            if template.source.format == 'xls':
                if not self.converter:
                    raise ERPError('spreadsheet_converter_unavailable', '请先将 XLS 另存为 XLSX。', 503)
                content = self.converter.convert(content)
                warnings.append('原 XLS 已转换为 XLSX，请核对转换后的版式。')
            with ZipFile(io.BytesIO(content)) as archive:
                if any(re.search(r'(^xl/(externalLinks|activeX|embeddings|slicers)/)|vbaProject', name, re.I)
                       for name in archive.namelist()):
                    raise ERPError('unsupported_workbook_objects', '模板含宏、外部链接或特殊对象，请提供普通 XLSX 空表。')
            workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = '逐季操作'
            sheet.append(['经营期', '操作', '数量', '金额'])
            count = report.snapshot.reading.problem.periods_per_year or 4
            for action in report.solution.plan.actions:
                sheet.append([period_label(action.period, count), action.description, action.quantity,
                              _excel_number(action.amount) if action.amount is not None else None])
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
        original_formulas = {(s.title, c.coordinate): c.value for s in workbook for row in s for c in row if c.data_type == 'f'}
        original_values = {(s.title, c.coordinate): c.value for s in workbook for row in s for c in row
                           if c.value is not None and c.data_type != 'f'}
        original_merges = {s.title: tuple(str(r) for r in s.merged_cells.ranges) for s in workbook}
        written = {}
        sources = output_sources(report.solution, report.snapshot)
        if mapping:
            for entry in mapping.entries:
                value = mapped_value(entry, sources)
                if value is None:
                    continue
                sheet = workbook[entry.sheet]
                cell = sheet[entry.cell]
                if (isinstance(cell, MergedCell) or cell.value is not None or cell.data_type == 'f'
                        or sheet.sheet_state != 'visible' or (sheet.protection.sheet and cell.protection.locked)
                        or sheet.row_dimensions[cell.row].hidden
                        or any(d.hidden and (d.min or 0) <= cell.column <= (d.max or 0)
                               for d in sheet.column_dimensions.values())):
                    raise ERPError('workbook_readonly', f'{entry.sheet} · {entry.cell} 不能覆盖。')
                if isinstance(value, Decimal):
                    cell.value = _excel_number(value)
                else:
                    cell.value = value
                    cell.data_type = 's'
                cell.comment = Comment(f'方案 {report.solution.id}\n来源：{", ".join(entry.source_ids)}\n{entry.explanation}', 'ERP Agent')
                written[(sheet.title, cell.coordinate)] = cell.value
        title = '填表说明'
        while title in workbook.sheetnames:
            title += '_'
        notes = workbook.create_sheet(title)
        notes.append(['方案编号', report.solution.id])
        notes.append(['题目版本', report.solution.problem_version])
        notes.append(['状态', '已验证' if report.solution.evaluation.verified else DRAFT_NOTICE])
        notes.append(['公式状态', '保留原公式；Excel 打开后自动计算' if original_formulas else '无公式'])
        notes.append(['表格', '单元格', '项目', '来源', '转换', '写入值', '说明'])
        if mapping:
            for entry in mapping.entries:
                value = mapped_value(entry, sources)
                notes.append([entry.sheet, entry.cell, entry.label, ', '.join(entry.source_ids), entry.conversion,
                              None if value is None else str(value), entry.explanation or ('未确定，保留原格' if value is None else '')])
        for row in notes:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for sheet in (notes,) if template else tuple(workbook):
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='3E58DB')
            for column in ('A', 'B', 'C', 'D', 'E', 'F', 'G'):
                sheet.column_dimensions[column].width = 22 if column != 'B' else 45
        workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        result = self.recalculator.recalculate(buffer.getvalue()) if original_formulas else RenderedArtifact(buffer.getvalue(), 'not_needed')
        if result.formula_status == 'recalculated':
            recalculated = load_workbook(io.BytesIO(result.content), data_only=False)
            try:
                changed = any(recalculated[s][c].value != value for (s, c), value in written.items())
                changed |= any(recalculated[s][c].value != value for (s, c), value in original_formulas.items())
                changed |= any(recalculated[s][c].value != value for (s, c), value in original_values.items())
                changed |= any(tuple(str(r) for r in recalculated[s].merged_cells.ranges) != ranges
                               for s, ranges in original_merges.items())
            finally:
                recalculated.close()
            if changed:
                result = RenderedArtifact(buffer.getvalue(), 'pending', ('重算改变了模板结构，已保留重算前文件。',))
        warnings.extend(result.warnings)
        reopened = load_workbook(io.BytesIO(result.content), data_only=False)
        try:
            if any(reopened[s][c].value != value for (s, c), value in written.items()):
                raise ERPError('workbook_verification', '导出表格写入值复查失败。', 500)
            if any(reopened[s][c].value != value for (s, c), value in original_formulas.items()):
                raise ERPError('workbook_verification', '导出表格公式发生变化。', 500)
            if any(reopened[s][c].value != value for (s, c), value in original_values.items()):
                raise ERPError('workbook_verification', '导出表格已有数据发生变化。', 500)
            if any(tuple(str(r) for r in reopened[s].merged_cells.ranges) != ranges for s, ranges in original_merges.items()):
                raise ERPError('workbook_verification', '导出表格合并区域发生变化。', 500)
        finally:
            reopened.close()
        return RenderedArtifact(result.content, result.formula_status, tuple(warnings))


def _excel_number(value: Decimal) -> int | float:
    if not value.is_finite() or len(value.as_tuple().digits) > 15 or abs(value) > Decimal('1e308'):
        raise ERPError('excel_precision', '数值超过 Excel 可安全表示的精度，请检查方案数据。')
    result = int(value) if value == value.to_integral_value() else float(value)
    if not math.isfinite(result) or (result == 0 and value != 0):
        raise ERPError('excel_precision', '数值超过 Excel 表示范围。')
    return result
