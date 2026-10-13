import io
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from erp_agent.adapters.spreadsheet_documents import XLSDocumentParser
from erp_agent.domain.errors import ERPError
from erp_agent.domain.models import (
    BudgetCell,
    BudgetFormulaReference,
    BudgetSheet,
    BudgetTemplate,
    BudgetValidation,
    ParsedDocument,
    SourceDocument,
)


class SpreadsheetBudgetParser:
    def parse(self, source: SourceDocument, content: bytes) -> BudgetTemplate:
        if source.format == "xls":
            document = XLSDocumentParser().parse(source, content)
            return self._from_xls(document)
        if source.format != "xlsx":
            raise ERPError("unsupported_template", "预算空表支持 XLSX 或 XLS。")
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 3000 or sum(e.file_size for e in entries) > 80 * 1024 * 1024:
                    raise ERPError("template_limit", "预算表解压体积或成员数过大。")
            workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
            cached = load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
        except (
            BadZipFile,
            InvalidFileException,
            ValueError,
            KeyError,
            OSError,
            SyntaxError,
        ) as exc:
            raise ERPError(
                "invalid_template", "预算表无法解析，请检查文件是否损坏或加密。"
            ) from exc
        sheets = []
        size = text_size = cell_count = 0
        warnings = ["只填写空白输入格，已有公式及数据保持原样；公式结果由 Excel 打开后自动计算。"]
        try:
            if len(workbook.worksheets) > 100:
                raise ERPError("template_limit", "预算表工作表过多。")
            for sheet in workbook.worksheets:
                size += sheet.max_row * sheet.max_column
                if size > 100000:
                    raise ERPError(
                        "template_limit", "预算表范围过大，请删除远处空白格式或拆分上传。"
                    )
                cells = []
                references = []
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell, MergedCell) or (
                            cell.value is None and not cell.has_style
                        ):
                            continue
                        text = "" if cell.value is None else str(cell.value)
                        text_size += len(text)
                        cell_count += 1
                        if text_size > 800000 or cell_count > 15000:
                            raise ERPError("template_limit", "预算表条目过多，请拆分上传。")
                        cells.append(
                            BudgetCell(
                                cell.coordinate,
                                text,
                                cell.number_format,
                                cell.data_type == "f",
                                not (sheet.protection.sheet and cell.protection.locked)
                                and not sheet.row_dimensions[cell.row].hidden
                                and not any(
                                    d.hidden and (d.min or 0) <= cell.column <= (d.max or 0)
                                    for d in sheet.column_dimensions.values()
                                ),
                                cell.alignment.horizontal or "",
                                cell.alignment.vertical or "",
                                bool(cell.alignment.wrap_text),
                                str(cached[sheet.title][cell.coordinate].value)
                                if cell.data_type == "f"
                                and cached[sheet.title][cell.coordinate].value is not None
                                else "",
                            )
                        )
                        if cell.data_type == "f":
                            references.extend(
                                _formula_references(cell.coordinate, text, sheet.title)
                            )
                validations = tuple(
                    BudgetValidation(
                        str(v.sqref),
                        v.type or "custom",
                        str(v.formula1 or ""),
                        str(v.formula2 or ""),
                    )
                    for v in sheet.data_validations.dataValidation
                )
                sheets.append(
                    BudgetSheet(
                        sheet.title,
                        sheet.max_row,
                        sheet.max_column,
                        tuple(cells),
                        tuple(str(r) for r in sheet.merged_cells.ranges),
                        sheet.sheet_state != "visible",
                        bool(sheet.protection.sheet),
                        tuple(references),
                        validations,
                        tuple(
                            _column_width(sheet, column)
                            for column in range(1, sheet.max_column + 1)
                        ),
                        tuple(
                            min(
                                1000,
                                max(
                                    24,
                                    int(
                                        (
                                            sheet.row_dimensions[row].height
                                            or sheet.sheet_format.defaultRowHeight
                                            or 15
                                        )
                                        * 4
                                        / 3
                                    ),
                                ),
                            )
                            for row in range(1, sheet.max_row + 1)
                        ),
                    )
                )
            if not any(c.text.strip() for s in sheets for c in s.cells):
                raise ERPError(
                    "empty_template", "预算表没有可识别的标题或条目，请提供带表头的空表。"
                )
            if any(s.hidden or s.protected for s in sheets):
                warnings.append("模板包含隐藏或保护工作表；不能自行覆盖受保护区域。")
        finally:
            workbook.close()
            cached.close()
        return BudgetTemplate(source, tuple(sheets), tuple(warnings))

    @staticmethod
    def _from_xls(document: ParsedDocument) -> BudgetTemplate:
        if (
            sum(len(b.cells) for b in document.blocks) > 15000
            or sum(len(c.text) for b in document.blocks for c in b.cells) > 800000
        ):
            raise ERPError("template_limit", "预算表条目过多，请拆分上传。")
        sheets = []
        for heading, table in zip(document.blocks[::2], document.blocks[1::2], strict=True):
            cells = tuple(
                BudgetCell(
                    f"{get_column_letter(c.column + 1)}{c.row + 1}",
                    c.text,
                    c.number_format,
                    writable=c.vertical_merge != "continue",
                )
                for c in table.cells
            )
            positions = {(c.row, c.column): c for c in table.cells}
            merges = []
            for cell in table.cells:
                if cell.vertical_merge == "continue":
                    continue
                if cell.column_span > 1 or cell.vertical_merge == "restart":
                    end_row = cell.row
                    while (
                        continuation := positions.get((end_row + 1, cell.column))
                    ) and continuation.vertical_merge == "continue":
                        end_row += 1
                    merges.append(
                        f"{get_column_letter(cell.column + 1)}{cell.row + 1}:"
                        f"{get_column_letter(cell.column + cell.column_span)}{end_row + 1}"
                    )
            sheets.append(
                BudgetSheet(
                    heading.text.removeprefix("工作表：").removesuffix("（隐藏）"),
                    max((c.row for c in table.cells), default=0) + 1,
                    max((c.column + c.column_span for c in table.cells), default=1),
                    cells,
                    tuple(merges),
                    hidden=heading.text.endswith("（隐藏）"),
                )
            )
        return BudgetTemplate(
            document.source,
            tuple(sheets),
            document.warnings
            + ("XLS 使用文件中已保存的结果；公式及保护信息未完整提取，正式填表需重新分析模板。",),
        )


def _formula_references(cell: str, formula: str, sheet: str) -> tuple[BudgetFormulaReference, ...]:
    try:
        tokens = Tokenizer(formula).items
    except TokenizerError:
        return ()
    references = []
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        target = sheet
        address = token.value
        if "!" in address:
            target, address = address.rsplit("!", 1)
            target = target.strip("'").replace("''", "'")
        address = address.replace("$", "")
        if "[" in target or not re.fullmatch(
            r"[A-Z]{1,3}[1-9][0-9]{0,6}(:[A-Z]{1,3}[1-9][0-9]{0,6})?", address
        ):
            continue
        reference = BudgetFormulaReference(cell, target, address)
        if reference not in references:
            references.append(reference)
    return tuple(references)


def _column_width(sheet, column: int) -> int:
    dimension = next(
        (d for d in sheet.column_dimensions.values() if (d.min or 0) <= column <= (d.max or 0)),
        None,
    )
    width = dimension.width if dimension else sheet.sheet_format.defaultColWidth or 8.43
    return min(2000, max(36, int(width * 7 + 5)))
