import io
from collections.abc import Sequence

from docx import Document
from docx.document import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table

from erp_agent.domain.errors import ERPError
from erp_agent.domain.models import Action, RenderedArtifact, ReportData
from erp_agent.domain.reporting import ReportNote, validate_report_outline

DRAFT_NOTICE = "方案草稿：经营数值尚未经过产能、现金、履约和财务仿真核算。"


def period_label(period: int, periods_per_year: int) -> str:
    year = (period - 1) // periods_per_year + 1
    quarter = (period - 1) % periods_per_year + 1
    return f"第 {year} 年 · 第 {quarter} 季"


class DocxReportRenderer:
    def render(self, report: ReportData) -> RenderedArtifact:
        solution, problem = report.solution, report.snapshot.reading.problem
        if report.outline:
            errors = validate_report_outline(
                report.outline,
                solution.plan,
                problem,
                report.snapshot.reading.rules,
            )
            if errors:
                raise ERPError("report_references", "方案文档引用检查失败：" + "；".join(errors))
        document = self._document(report)
        count = problem.periods_per_year or 4
        horizon = (problem.years or 4) * count
        document.add_heading(solution.plan.title, 0)
        document.add_paragraph(problem.title)
        if not solution.evaluation.verified:
            document.add_paragraph(DRAFT_NOTICE)
        document.add_paragraph(
            f"方案编号：{solution.id}；题目版本：{solution.problem_version}；模型：{solution.model}"
        )
        document.add_paragraph(
            f"经营周期：{problem.years or 4} 年，每年 {count} 季，共 {horizon} 期。"
        )
        self._parameters(document, report)
        self._milestones(document, report, count)
        self._strategy(document, report, count)
        self._operation_cards(document, report, count, horizon)
        self._budgets(document, report, count)
        self._risks(document, report, count)
        self._closing(document, report)
        self._sources(document, report)
        content = io.BytesIO()
        document.save(content)
        reopened = Document(io.BytesIO(content.getvalue()))
        headings = {p.text for p in reopened.paragraphs}
        if any(period_label(p, count) not in headings for p in range(1, horizon + 1)):
            raise ERPError("report_verification", "方案文档周期检查失败。", 500)
        return RenderedArtifact(
            content.getvalue(),
            warnings=() if solution.evaluation.verified else (DRAFT_NOTICE,),
        )

    @staticmethod
    def _document(report: ReportData) -> WordDocument:
        document = Document()
        section = document.sections[0]
        section.header.paragraphs[0].text = "ERP 经营方案 · " + (
            "已验证" if report.solution.evaluation.verified else "草稿"
        )
        normal = document.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        normal.paragraph_format.space_after = Pt(6)
        section.top_margin = section.bottom_margin = Inches(0.75)
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            document.styles[name].font.color.rgb = RGBColor.from_string("3E58DB")
        footer = section.footer.paragraphs[0]
        footer.text = f"方案 {report.solution.id} · 题目版本 {report.solution.problem_version}    "
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        return document

    @staticmethod
    def _table(document: WordDocument, headers: Sequence[str]) -> Table:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Light Shading Accent 1"
        for cell, text in zip(table.rows[0].cells, headers, strict=True):
            cell.text = text
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        return table

    def _parameters(self, document: WordDocument, report: ReportData) -> None:
        document.add_heading("一、题目参数速查", 1)
        document.add_paragraph("先核对条件，再看战略和操作；具体数值与来源编号沿用已核对题目。")
        facts = [
            fact
            for fact in report.snapshot.reading.problem.facts
            if fact.category not in {"order_pool", "order_contract", "demand_forecast"}
        ]
        if not facts:
            document.add_paragraph("未登记基础参数，不能从输出样例补入。")
            return
        table = self._table(document, ("对象 / 参数", "值与单位", "状态 / 引用"))
        for fact in facts:
            row = table.add_row().cells
            row[0].text = f"{fact.entity_name} · {fact.key}"
            row[1].text = "未确定" if fact.value is None else f"{fact.value} {fact.unit}".strip()
            row[2].text = f"{self._state(fact.state.value)} · {fact.id}"

    def _strategy(self, document: WordDocument, report: ReportData, count: int) -> None:
        document.add_heading("三、战略路线与取舍", 1)
        document.add_paragraph(report.solution.plan.summary)
        tones = {
            "conservative": "保守",
            "balanced": "平稳",
            "aggressive": "激进",
            "custom": "自定义",
        }
        config = report.solution.generation
        document.add_paragraph("方案基调：" + tones[config.tone.value])
        if config.custom_prompt:
            document.add_paragraph("自定义要求：" + config.custom_prompt)
        for requirement in config.requirements:
            document.add_paragraph("调整要求：" + requirement)
        if report.outline and report.outline.strategy:
            self._notes(document, report, report.outline.strategy, count)
        else:
            document.add_paragraph("按下列逐期行动执行；当前方案未附额外的战略解释。")
        document.add_heading("默认与暂定条件", 2)
        self._list(document, report.solution.plan.assumptions, "无已登记假设。")

    def _milestones(self, document: WordDocument, report: ReportData, count: int) -> None:
        document.add_heading("二、关键时点与前置安排", 1)
        document.add_paragraph(
            "对照接单与交付窗口，核对市场、认证、研发、安装、材料和人员的准备时点。"
            "这里引用当前计划的安排，未核实的提前期不自动补算。"
        )
        if report.outline and report.outline.milestones:
            self._notes(document, report, report.outline.milestones, count)
        else:
            document.add_paragraph("当前方案未登记带依据的时点说明，操作日期见逐期卡。")

    @staticmethod
    def _action_group(action: Action, report: ReportData) -> str:
        kind = action.kind.lower()
        groups = (
            (
                "财务",
                "financ loan repay interest discount cash reserve tax depreciat wage "
                "salary pay budget 财务 融资 借款 还款 利息 贴现 现金 税 折旧 工资 费用 预算",
            ),
            (
                "人力",
                "hire recruit worker train incent shift fire dismiss "
                "人力 招工 招聘 培训 激励 班次 解雇",
            ),
            (
                "生产",
                "produc manufact material procure purchase line install maintain inventory "
                "生产 采购 原料 产线 安装 维修 库存",
            ),
            (
                "营销",
                "market certif research develop advert bid sale order deliver "
                "营销 市场 认证 研发 广告 竞单 接单 交付",
            ),
        )
        for title, tokens in groups:
            if any(token in kind for token in tokens.split()):
                return title
        categories = {
            fact.category
            for fact in report.snapshot.reading.problem.facts
            if fact.entity_id == action.entity_id
        }
        for title, values in (
            ("营销", {"market", "certification", "product", "order_pool", "order_contract"}),
            ("生产", {"production_line", "material", "bom"}),
            ("人力", {"worker", "training", "incentive", "shift"}),
            ("财务", {"finance", "discount", "budget"}),
        ):
            if categories & values:
                return title
        return "其他操作"

    def _operation_cards(
        self,
        document: WordDocument,
        report: ReportData,
        count: int,
        horizon: int,
    ) -> None:
        document.add_heading("四、逐季经营操作卡", 1)
        document.add_paragraph(
            "按行动编号和前置条件确定执行顺序；营销、生产、人力、财务分类仅供查阅和核对。"
            "行动编号与当前方案一致，金额和数量直接引用计划。"
        )
        names = {fact.entity_id: fact.entity_name for fact in report.snapshot.reading.problem.facts}
        for period in range(1, horizon + 1):
            document.add_heading(period_label(period, count), 2)
            actions = [
                (index, action)
                for index, action in enumerate(report.solution.plan.actions)
                if action.period == period
            ]
            if not actions:
                document.add_paragraph("此期缺少操作，需完善计划；不能视为无需经营或已经收官。")
            for group in ("营销", "生产", "人力", "财务", "其他操作"):
                selected = [
                    (index, action)
                    for index, action in actions
                    if self._action_group(action, report) == group
                ]
                if group == "其他操作" and not selected:
                    continue
                document.add_heading(group, 3)
                if not selected:
                    document.add_paragraph("本方案未登记此类操作。")
                    continue
                table = self._table(document, ("对象", "操作与前置条件", "数量", "金额"))
                for _index, action in selected:
                    row = table.add_row().cells
                    row[0].text = names.get(action.entity_id, action.entity_id)
                    row[1].text = f"行动 {action.action_id} · {action.kind}\n{action.description}"
                    if action.depends_on:
                        row[1].text += "\n前置行动：" + "、".join(action.depends_on)
                    if action.completion_period is not None:
                        row[1].text += "\n完成：" + period_label(action.completion_period, count)
                    if action.payment_period is not None:
                        row[1].text += "\n支付 / 收款：" + period_label(
                            action.payment_period, count
                        )
                    if action.missing_reason:
                        row[1].text += "\n未确定：" + action.missing_reason
                    row[2].text = (
                        f"{action.quantity} {action.quantity_unit}".strip()
                        if action.quantity is not None
                        else "未确定 / 不适用"
                    )
                    row[3].text = (
                        f"{action.amount} {action.amount_unit} {action.amount_meaning}".strip()
                        if action.amount is not None
                        else "未确定 / 不适用"
                    )
            if report.outline:
                notes = tuple(
                    note
                    for note in (*report.outline.milestones, *report.outline.risks)
                    if note.period == period
                )
                if notes:
                    document.add_heading("本季条件与风险", 3)
                    self._notes(document, report, notes, count, level=4)
            document.add_paragraph(
                "执行核对：以赛题事件顺序检查资格、到货与完工、可用产能、付款时点和交付期限。"
            )
            if period % count == 0:
                document.add_paragraph(
                    "年末核对：按适用规则结算工资、费用、折旧、税费与债务，核对库存和报表衔接。"
                )
            if period == horizon:
                document.add_paragraph(
                    "经营终点：核对尚未交付订单、在制品、库存、应收应付、债务和税费的最终处理。"
                )

    def _budgets(self, document: WordDocument, report: ReportData, count: int) -> None:
        document.add_heading("五、预算规划与填表依据", 1)
        items = report.solution.plan.budget_items
        if not items:
            document.add_paragraph("未附模板预算条目。操作值见逐季计划；不据此虚构财务报表。")
            return
        table = self._table(document, ("项目与位置", "期间", "草稿值", "依据 / 待确定原因"))
        for index, item in enumerate(items):
            row = table.add_row().cells
            row[0].text = f"预算 B{index + 1} · {item.label}\n{item.sheet} · {item.cell}"
            row[1].text = period_label(item.period, count) if item.period else "汇总 / 未确定"
            row[2].text = "未确定" if item.value is None else f"{item.value} {item.unit}".strip()
            actions = "、".join(item.action_ids or tuple(f"A{i + 1}" for i in item.action_indices))
            row[3].text = item.explanation + (f"\n关联行动：{actions}" if actions else "")

    def _risks(self, document: WordDocument, report: ReportData, count: int) -> None:
        document.add_heading("六、风险、触发条件与应对", 1)
        if report.outline and report.outline.risks:
            self._notes(document, report, report.outline.risks, count)
        document.add_heading("当前方案登记风险", 2)
        self._list(document, report.solution.plan.risks, "无已登记风险；不代表经营没有风险。")
        document.add_paragraph(
            "备选响应仅解释当前计划及其适用条件；新增操作或改变预算须生成新的方案版本。"
        )

    def _closing(self, document: WordDocument, report: ReportData) -> None:
        document.add_heading("七、收官与待处理事项", 1)
        evaluation = report.solution.evaluation
        document.add_paragraph("核验范围：" + evaluation.scope)
        if evaluation.verified:
            document.add_paragraph("经营校验已通过。")
        else:
            document.add_paragraph(DRAFT_NOTICE)
        if evaluation.score is not None:
            document.add_paragraph("当前求解结果得分：" + str(evaluation.score))
        else:
            document.add_paragraph("当前结果未提供可核算的最终得分。")
        self._list(
            document, tuple(issue.message for issue in evaluation.issues), "无已登记校验问题。"
        )
        reading_issues = tuple(issue.message for issue in report.snapshot.reading.issues)
        if reading_issues:
            document.add_heading("题目仍登记的事项", 2)
            self._list(document, reading_issues, "")

    def _notes(
        self,
        document: WordDocument,
        report: ReportData,
        notes: tuple[ReportNote, ...],
        count: int,
        level: int = 2,
    ) -> None:
        facts = {fact.id: fact for fact in report.snapshot.reading.problem.facts}
        rules = {rule.id: rule for rule in report.snapshot.reading.rules.rules}
        names = {fact.entity_id: fact.entity_name for fact in facts.values()}
        for note in notes:
            document.add_heading(note.title, level)
            if note.period:
                document.add_paragraph("适用时点：" + period_label(note.period, count))
            document.add_paragraph(note.explanation)
            if note.condition:
                document.add_paragraph("触发条件：" + note.condition)
            if note.response:
                document.add_paragraph("应对方式：" + note.response)
            table = self._table(document, ("依据", "当前方案 / 题目中的内容"))
            identifiers = note.action_ids or tuple(
                report.solution.plan.actions[i].action_id for i in note.action_indices
            )
            for identifier in identifiers:
                index, action = next(
                    (i, a)
                    for i, a in enumerate(report.solution.plan.actions)
                    if a.action_id == identifier
                )
                row = table.add_row().cells
                row[0].text = f"行动 {action.action_id} · {period_label(action.period, count)}"
                row[
                    1
                ].text = f"{names.get(action.entity_id, action.entity_id)} · {action.description}"
                if action.quantity is not None:
                    row[1].text += f"；数量 {action.quantity}"
                if action.amount is not None:
                    row[1].text += f"；金额 {action.amount}"
            for fact_id in note.fact_ids:
                fact = facts[fact_id]
                row = table.add_row().cells
                row[0].text = f"参数 {fact.id} · {self._state(fact.state.value)}"
                row[1].text = f"{fact.entity_name} · {fact.key}：" + (
                    "未确定" if fact.value is None else f"{fact.value} {fact.unit}".strip()
                )
            for rule_id in note.rule_ids:
                rule = rules[rule_id]
                row = table.add_row().cells
                row[0].text = f"规则 {rule.id} · {self._state(rule.state.value)}"
                row[1].text = f"{rule.name}：{rule.expression}"
            for index in note.budget_indices:
                item = report.solution.plan.budget_items[index]
                row = table.add_row().cells
                row[0].text = f"预算 B{index + 1} · {item.sheet}!{item.cell}"
                row[1].text = (
                    item.label
                    + "："
                    + ("未确定" if item.value is None else f"{item.value} {item.unit}".strip())
                )

    def _sources(self, document: WordDocument, report: ReportData) -> None:
        document.add_heading("附录：完整题目数据与原文依据", 1)
        evidence = {source.id: source for source in report.sources}
        table = self._table(document, ("对象 / 字段", "值与单位", "状态", "原文依据 / 说明"))
        for fact in report.snapshot.reading.problem.facts:
            row = table.add_row().cells
            row[0].text = f"{fact.entity_name} · {fact.key}\n参数 {fact.id}"
            row[1].text = "未确定" if fact.value is None else f"{fact.value} {fact.unit}".strip()
            row[2].text = self._state(fact.state.value)
            references = [
                f"{evidence[e].location}：{evidence[e].text}"
                for e in fact.evidence_ids
                if e in evidence
            ]
            row[3].text = "\n".join(references + ([fact.explanation] if fact.explanation else []))
        document.add_heading("适用规则", 1)
        for rule in report.snapshot.reading.rules.rules:
            document.add_heading(rule.name, 2)
            document.add_paragraph(f"规则 {rule.id} · {self._state(rule.state.value)}")
            document.add_paragraph(rule.expression)
            if rule.explanation:
                document.add_paragraph(rule.explanation)
            for source_id in rule.evidence_ids:
                if source_id in evidence:
                    source = evidence[source_id]
                    document.add_paragraph(f"原文依据：{source.location}：{source.text}")

    @staticmethod
    def _state(value: str) -> str:
        return {
            "known": "已知",
            "assumed": "默认 / 假设",
            "conflicted": "存在冲突",
            "missing": "未提供",
        }[value]

    @staticmethod
    def _list(document: WordDocument, values: Sequence[str], empty: str) -> None:
        for value in values:
            document.add_paragraph(value, style="List Bullet")
        if not values and empty:
            document.add_paragraph(empty)
