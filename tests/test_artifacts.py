import io
import unittest
from unittest.mock import patch
from dataclasses import replace
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from docx import Document

from erp_agent.adapters.artifacts import DocxReportRenderer, XlsxWorkbookRenderer
from erp_agent.adapters.budget_templates import SpreadsheetBudgetParser
from erp_agent.application.budget_context import budget_material
from erp_agent.application.codec import budget_template
from erp_agent.application.contracts import decode_proposal, dumps, loads, to_json
from erp_agent.application.workbook_mapper import WorkbookMapper
from erp_agent.domain.models import BudgetItem, ReportData, SolutionBundle, SolverMode, SourceDocument, WorkbookEntry
from erp_agent.domain.errors import ERPError
from erp_agent.domain.reporting import ReportNote, ReportOutline
from erp_agent.domain.workbook import validate_mapping
from tests.helpers import snapshot, proposal_json, Control, FakeLLM
from erp_agent.application.solver import DraftPlanEvaluator


def report():
    problem = snapshot()
    plan = decode_proposal(proposal_json(), 'fake').baseline
    solution = SolutionBundle('solution', problem.project_id, 1, SolverMode.LLM_ONLY,
                              plan, DraftPlanEvaluator().evaluate(problem, plan), None, 1, 1, None, 'fake', '')
    return ReportData(problem, solution)


def template_fixture(transposed=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '借款预算' if not transposed else '年度预算'
    if transposed:
        sheet.append(['期间', '借款'])
        sheet.append(['第1季', None])
        sheet['B2'].number_format = '0.00'
    else:
        sheet.append(['项目', '第1季', '第2季'])
        sheet.append(['借款', None, None])
        sheet['B2'].number_format = '0.00'
        sheet['B3'] = '=SUM(B2)'
        sheet.merge_cells('A4:C4')
        sheet['A4'] = '说明'
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()
    original = content.getvalue()
    template = SpreadsheetBudgetParser().parse(SourceDocument('template', 'budget.xlsx', 'hash', len(original), 'xlsx'), original)
    return original, template


class ArtifactTests(unittest.TestCase):
    def test_formula_driven_inputs_reuse_fixed_budget_without_model_requests(self):
        data = report()
        workbook = Workbook()
        first = workbook.active
        first.title = '第一年'
        first.append(['项目', '第一季', '年度合计'])
        first.append(['借款', None, '=SUM(B2)'])
        first['B2'].number_format = '0.00'
        first['B3'] = 0
        second = workbook.create_sheet('生产')
        second.append(['期间', '操作', '期初资本', '跨表合计'])
        second.append(['第一季', None, None, "='第一年'!C2"])
        second['B2'].number_format = '@'
        second['C2'].number_format = '0'
        second['A3'] = '内部'
        second['B3'].number_format = '0'
        second.row_dimensions[3].hidden = True
        validation = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=0)
        validation.add(second['C2'])
        second.add_data_validation(validation)
        buffer = io.BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
        template = SpreadsheetBudgetParser().parse(SourceDocument('template', 'budget.xlsx', 'hash', len(content), 'xlsx'), content)
        self.assertEqual(budget_template(to_json(template)), template)
        material = budget_material((template,), include_formula_references=True)[0]
        self.assertEqual(material['sheets'][0]['formula_references'], [['C2', '第一年', ['B2']]])
        self.assertEqual(material['sheets'][1]['validations'][0]['kind'], 'whole')
        self.assertFalse(next(c for c in template.sheets[1].cells if c.coordinate == 'B3').writable)
        old = BudgetItem('template', '第一年', 'B2', '借款', 1, '元', '10', (0,), '旧预算')
        data = replace(data, solution=replace(data.solution, plan=replace(data.solution.plan, budget_items=(
            old,
            BudgetItem('template', '生产', 'B2', '操作', 1, '', data.solution.plan.actions[0].description, (0,), '方案操作'),
            BudgetItem('template', '生产', 'C2', '期初资本', None, '元', '100000', (), '题目初始资本'),
        ))))

        class MappingLLM(FakeLLM):
            def generate(self, request):
                current = loads(request.user)['template']['sheets'][0]['name']
                entries = [
                    {'sheet': current, 'cell': 'B2', 'label': '借款' if current == '第一年' else '操作',
                     'source_ids': ['action:0:amount' if current == '第一年' else 'action:0:description'],
                     'conversion': 'identity', 'explanation': '第一季方案'},
                    {'sheet': current, 'cell': 'C2', 'label': '汇总' if current == '第一年' else '期初资本',
                     'source_ids': [] if current == '第一年' else ['fact:cash'],
                     'conversion': 'identity', 'explanation': '保留公式或引用题目初始值'},
                ]
                self.response = dumps({'entries': entries})
                return super().generate(request)

        llm = MappingLLM('')
        with patch('erp_agent.application.workbook_mapper.MAPPING_BATCH_CHARS', 1):
            mapping = WorkbookMapper().analyze(data.solution, template, llm, Control(), data.snapshot)
        self.assertEqual(len(llm.calls), 0)
        self.assertEqual(mapping.entries[0].source_ids, ('budget:0',))
        self.assertEqual(len(mapping.entries), 3)
        self.assertFalse(any(i.severity == 'error' for i in mapping.issues))
        filled = XlsxWorkbookRenderer().render(data, template, content, replace(mapping, confirmed=True))
        reopened = load_workbook(io.BytesIO(filled.content), data_only=False)
        self.assertEqual(reopened['第一年']['C2'].value, '=SUM(B2)')
        self.assertEqual(reopened['第一年']['B3'].value, 0)
        self.assertEqual(reopened['生产']['D2'].value, "='第一年'!C2")
        self.assertEqual(reopened['生产']['B2'].value, data.solution.plan.actions[0].description)
        self.assertEqual(reopened['生产']['C2'].value, 100000)
        self.assertTrue(reopened['生产'].row_dimensions[3].hidden)
        self.assertEqual(str(reopened['生产'].data_validations), str(second.data_validations))
        self.assertTrue(any(i.severity == 'error' for i in validate_mapping(
            template, (WorkbookEntry('生产', 'B3', '内部', ('action:0:amount',)),), data.solution)))
        changed = tuple(replace(e, source_ids=('action:1:amount',)) if e.cell == 'B2' and e.sheet == '第一年'
                        else e for e in mapping.entries)
        self.assertIn('mapping.plan_mismatch', {i.code for i in validate_mapping(template, changed, data.solution)})
        self.assertIn('mapping.plan_omitted', {i.code for i in validate_mapping(template, mapping.entries[1:], data.solution)})
        reopened.close()
        workbook.close()

    def test_report_covers_all_sixteen_periods_and_preserves_operating_values(self):
        data = report()
        problem = replace(data.snapshot, reading=replace(data.snapshot.reading, problem=replace(data.snapshot.reading.problem, years=4, periods_per_year=4)))
        plan = replace(data.solution.plan, actions=tuple(replace(data.solution.plan.actions[0], period=p) for p in range(1, 17)))
        result = DocxReportRenderer().render(replace(data, snapshot=problem, solution=replace(data.solution, plan=plan)))
        document = Document(io.BytesIO(result.content))
        self.assertIn('第 4 年 · 第 4 季', [p.text for p in document.paragraphs])
        tables = [t for t in document.tables if t.cell(0, 0).text == '对象']
        self.assertEqual(len(tables), 16)
        self.assertTrue(all(t.cell(1, 3).text == '10' for t in tables))
        self.assertTrue(any('草稿' in p.text for p in document.paragraphs))

    def test_report_uses_actual_cycle_and_keeps_final_quarter_after_sales_end(self):
        data = report()
        problem = replace(data.snapshot.reading.problem, years=2, periods_per_year=3)
        data = replace(data, snapshot=replace(data.snapshot,
                       reading=replace(data.snapshot.reading, problem=problem)))
        document = Document(io.BytesIO(DocxReportRenderer().render(data).content))
        headings = [p.text for p in document.paragraphs]
        quarters = [text for text in headings if text.startswith('第 ') and ' · 第 ' in text]
        self.assertEqual(len(quarters), 6)
        self.assertIn('第 2 年 · 第 3 季', quarters)
        self.assertNotIn('第 4 年 · 第 4 季', quarters)
        self.assertTrue(any('不能视为无需经营或已经收官' in text for text in headings))
        chapters = [text for text in headings if text.startswith(('一、', '二、', '三、', '四、'))]
        self.assertEqual(chapters, ['一、题目参数速查', '二、关键时点与前置安排',
                                    '三、战略路线与取舍', '四、逐季经营操作卡'])

    def test_collaborative_notes_insert_referenced_plan_and_fact_values(self):
        data = report()
        action = replace(data.solution.plan.actions[0], amount=Decimal('10.50'), quantity=7)
        budget = BudgetItem('template', '预算', 'B2', '运营准备金', 1, '元', '10.50', (0,), '沿用行动金额')
        plan = replace(data.solution.plan, actions=(action, *data.solution.plan.actions[1:]),
                       budget_items=(budget,))
        note = ReportNote('先保障资金再执行', '根据当前行动安排付款。', (0,), ('cash',),
                          budget_indices=(0,), period=1, condition='付款前资金不足',
                          response='复核当前计划的付款安排')
        data = replace(data, solution=replace(data.solution, plan=plan),
                       outline=ReportOutline(strategy=(note,), milestones=(note,), risks=(note,)))
        document = Document(io.BytesIO(DocxReportRenderer().render(data).content))
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('触发条件：付款前资金不足', text)
        cells = '\n'.join(cell.text for table in document.tables
                          for row in table.rows for cell in row.cells)
        self.assertIn('数量 7；金额 10.50', cells)
        self.assertIn('公司 · initial_cash：100000 元', cells)
        self.assertIn('运营准备金：10.50 元', cells)

    def test_report_rejects_notes_from_another_solution_or_problem(self):
        data = report()
        for note in (ReportNote('战略', '说明', action_indices=(99,)),
                     ReportNote('战略', '说明', fact_ids=('other-project-fact',)),
                     ReportNote('战略', '说明', action_indices=(0,), period=3)):
            with self.subTest(note=note), self.assertRaises(ERPError) as caught:
                DocxReportRenderer().render(replace(data, outline=ReportOutline(strategy=(note,))))
            self.assertEqual(caught.exception.code, 'report_references')

    def test_two_template_layouts_reference_sources_without_inventing_values(self):
        data = report()
        for transposed in (False, True):
            original, template = template_fixture(transposed)
            sheet_name = template.sheets[0].name
            entry = {'sheet': sheet_name, 'cell': 'B2', 'label': '借款', 'source_ids': ['action:0:amount'], 'conversion': 'identity', 'explanation': '第一季借款操作金额'}
            llm = FakeLLM(dumps({'entries': [entry, {**entry, 'cell': 'B3' if not transposed else 'A2', 'source_ids': [], 'explanation': '保留原值或公式'}]}))
            mapping = WorkbookMapper().analyze(data.solution, template, llm, Control())
            self.assertFalse(any(i.severity == 'error' for i in mapping.issues))
            filled = XlsxWorkbookRenderer().render(data, template, original, replace(mapping, confirmed=True))
            reopened = load_workbook(io.BytesIO(filled.content), data_only=False)
            self.assertEqual(reopened[sheet_name]['B2'].value, 10)
            self.assertEqual(reopened[sheet_name]['B2'].number_format, '0.00')
            if not transposed:
                self.assertEqual(reopened[sheet_name]['B3'].value, '=SUM(B2)')
                self.assertIn('A4:C4', [str(r) for r in reopened[sheet_name].merged_cells.ranges])
                self.assertEqual(filled.formula_status, 'pending')
            self.assertTrue(any('草稿' in str(c.value) for row in reopened['填表说明'] for c in row))
            reopened.close()
            unchanged = load_workbook(io.BytesIO(original))
            self.assertIsNone(unchanged[sheet_name]['B2'].value)
            unchanged.close()

    def test_readonly_and_unknown_sources_cannot_be_confirmed(self):
        data = report()
        _, template = template_fixture()
        sheet = template.sheets[0].name
        for entry in (WorkbookEntry(sheet, 'B3', '公式', ('action:0:amount',)),
                      WorkbookEntry(sheet, 'B4', '合并续格', ('action:0:amount',)),
                      WorkbookEntry(sheet, 'A2', '标题', ('action:0:amount',)),
                      WorkbookEntry(sheet, 'B2', '未知', ('fake-value',))):
            self.assertTrue(any(i.severity == 'error' for i in validate_mapping(template, (entry,), data.solution)))
