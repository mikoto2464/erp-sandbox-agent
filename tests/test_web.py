import tempfile
import io
import time
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path
from urllib.error import HTTPError
from types import SimpleNamespace

from erp_agent.api.server import Handler, create_access, create_workflow
from erp_agent.application.contracts import dumps, loads
from erp_agent.application.input_sources import text_source
from erp_agent.adapters.documents import OOXMLDocumentParser
from erp_agent.adapters.persistence import SQLiteRepository
from erp_agent.config import Settings
from erp_agent.domain.models import SourceDocument
from erp_agent.application.ports import LLMResponse
from unittest.mock import patch
from tests.helpers import FakeLLM, docx_bytes, proposal_json
from openpyxl import Workbook
from openpyxl import load_workbook
from docx import Document


class WebTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="erp-test-")
        self.llm = FakeLLM("")
        settings = Settings(port=8000, data_dir=Path(self.temp.name), access_key="test-access")
        workflow = create_workflow(settings, self.llm)
        self.server = SimpleNamespace(settings=settings, workflow=workflow,
                                      access=create_access(settings), models=workflow.model_settings)
        self.cookie = ""
        self.response_headers = ""
        self.request("/api/auth/login", {"key": "test-access", "remember": True})

    def tearDown(self):
        self.server.workflow.jobs.close()
        self.temp.cleanup()

    def request(self, path, data=None, content=None, headers=None):
        body = dumps(data).encode() if data is not None else content
        actual_headers = headers or ({"Content-Type": "application/json"} if data is not None else {})
        code, result = self.raw_request(path, body, actual_headers)
        if code >= 400:
            raise HTTPError(path, code, result.decode(), None, None)
        return loads(result.decode())

    def raw_request(self, path, body=None, headers=None, method=None):
        method = method or ("POST" if body is not None else "GET")
        content = body or b""
        lines = [f"{method} {path} HTTP/1.1", "Host: localhost:8000", "Connection: close",
                 f"Content-Length: {len(content)}"]
        lines.extend(f"{k}: {v}" for k, v in (headers or {}).items())
        if self.cookie:
            lines.append("Cookie: " + self.cookie)
        raw = ("\r\n".join(lines) + "\r\n\r\n").encode() + content

        class MemorySocket:
            def __init__(self):
                self.input = io.BytesIO(raw)
                self.output = bytearray()
            def makefile(self, mode, buffering=None):
                return self.input
            def sendall(self, data):
                self.output.extend(data)
            def settimeout(self, value):
                pass

        connection = MemorySocket()
        Handler(connection, ("127.0.0.1", 12345), self.server)
        head, result = bytes(connection.output).split(b"\r\n\r\n", 1)
        self.response_headers = head.decode()
        for line in self.response_headers.split("\r\n"):
            if line.startswith("Set-Cookie: "):
                self.cookie = line[len("Set-Cookie: "):].split(";", 1)[0]
        return int(head.split(b" ")[1]), result

    def wait_job(self, project, job):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            result = self.request(f"/api/projects/{project}/jobs/{job}")
            if result["status"] not in {"running", "queued"}:
                return result
            time.sleep(0.01)
        self.fail("后台任务没有完成")

    def setup_problem(self):
        project = self.request("/api/projects", {"name": "离线端到端"})["id"]
        base = f"/api/projects/{project}"
        job = self.request(base + "/documents", content=docx_bytes(), headers={
            "Content-Type": "application/octet-stream", "X-Filename": "topic.docx",
        })
        self.assertEqual(self.wait_job(project, job["id"])["status"], "succeeded")
        docs = self.request(base + "/documents")
        eid = docs[0]["evidence"][0]["id"]
        facts = []
        for key, value, kind in [("initial_cash", "100000", "number"),
                                 ("years", "1", "integer"), ("periods_per_year", "2", "integer")]:
            facts.append({"id": key, "category": "initial", "entity_id": "company",
                          "entity_name": "公司", "key": key, "value": value, "value_type": kind,
                          "unit": "", "state": "known", "evidence_ids": [eid],
                          "conversion": "identity", "explanation": ""})
        for fact in facts[1:]:
            fact["state"] = "assumed"
            fact["explanation"] = "测试情景：经营一年，每年两期"
        self.llm.response = dumps({"title": "测试题目", "facts": facts, "rules": [], "issues": []})
        job = self.request(base + "/read", {})
        self.assertEqual(self.wait_job(project, job["id"])["status"], "succeeded")
        result = self.request(base + "/problem")
        return project, result

    def test_budget_tone_and_iterative_solution_context(self):
        project, _ = self.setup_problem()
        base = f'/api/projects/{project}'
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '现金预算'
        sheet.append(['项目', '第1季', '第2季'])
        sheet.append(['借款', None, None])
        sheet['B2'].number_format = '0.00'
        sheet['B3'] = '=SUM(B2)'
        sheet.merge_cells('A4:C4')
        sheet['A4'] = '备注'
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        job = self.request(base + '/templates', content=content.getvalue(), headers={
            'Content-Type': 'application/octet-stream', 'X-Filename': 'budget.xlsx',
        })
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        template = self.request(base + '/templates')[0]
        self.assertEqual(template['sheets'][0]['merged_ranges'], ['A4:C4'])
        self.assertTrue(next(c for c in template['sheets'][0]['cells'] if c['coordinate'] == 'B3')['formula'])
        proposal = loads(proposal_json())
        proposal['plan']['budget_items'] = [{
            'template_id': template['source']['id'], 'sheet': '现金预算', 'cell': 'B2',
            'label': '借款', 'period': 1, 'unit': '元', 'value': '10', 'action_indices': [0],
            'explanation': '依据第一期操作安排，尚未完成仿真验证',
        }]
        self.llm.response = dumps(proposal)
        generation = {'tone': 'conservative', 'custom_prompt': '',
                      'template_ids': [template['source']['id']]}
        job = self.request(base + '/solve', {'version': 1, 'generation': generation})
        done = self.wait_job(project, job['id'])
        self.assertEqual(done['status'], 'succeeded')
        first = self.request(base + '/solutions')[0]
        payload = loads(self.llm.calls[-1].user)
        self.assertEqual(payload['budget_templates'][0]['source']['id'], template['source']['id'])
        self.assertIn('保守', payload['tone_instruction'])
        self.assertIsNone(payload['previous_plan'])
        self.assertEqual(first['generation']['tone'], 'conservative')
        self.assertEqual(job['metadata']['generation']['template_ids'], generation['template_ids'])
        old_raw = first['raw_response']
        for parent, requirement in [(first['id'], '第一年不增产线'), (None, '继续减少借款')]:
            parent = parent or self.request(base + '/solutions')[-1]['id']
            generation.update(previous_solution_id=parent, revision_prompt=requirement)
            proposal['plan']['title'] = requirement
            self.llm.response = dumps(proposal)
            job = self.request(base + '/solve', {'version': 1, 'generation': generation})
            self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        plans = self.request(base + '/solutions')
        self.assertEqual(len(plans), 3)
        self.assertEqual(plans[0]['raw_response'], old_raw)
        self.assertEqual(plans[-1]['generation']['requirements'], ['第一年不增产线', '继续减少借款'])
        payload = loads(self.llm.calls[-1].user)
        self.assertEqual(payload['previous_plan']['title'], '第一年不增产线')
        self.assertNotIn('raw_response', payload)
        self.assertEqual(len(self.request(base + '/documents')), 1)
        other = self.request('/api/projects', {'name': '其他'})['id']
        other_base = f'/api/projects/{other}'
        with self.assertRaises(HTTPError) as forbidden:
            self.request(other_base + f"/templates/{template['source']['id']}/retry", {})
        self.assertEqual(forbidden.exception.code, 404)
        for config, code in [({'tone': 'invalid'}, 400), ({'tone': 'custom'}, 400),
                             ({'previous_solution_id': first['id']}, 400),
                             ({'template_ids': ['not-in-project']}, 404)]:
            with self.subTest(config=config), self.assertRaises(HTTPError) as rejected:
                self.request(base + '/solve', {'version': 1, 'generation': config})
            self.assertEqual(rejected.exception.code, code)
        reading = self.request(base + '/problem')['snapshot']['reading']
        self.request(base + '/revisions', {'expected_version': 1, 'reading': {
            'title': '新版', 'facts': reading['problem']['facts'], 'rules': [], 'issues': [],
        }})
        with self.assertRaises(HTTPError) as stale:
            self.request(base + '/solve', {'version': 2, 'generation': generation})
        self.assertEqual(stale.exception.code, 409)

    def test_mapping_confirmation_export_download_and_project_isolation(self):
        project, _ = self.setup_problem()
        base = f'/api/projects/{project}'
        self.llm.response = proposal_json()
        job = self.request(base + '/solve', {'version': 1})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        solution = self.request(base + '/solutions')[0]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '资金计划'
        sheet.append(['项目', '第一季'])
        sheet.append(['借款', None])
        sheet['B2'].number_format = '0.00'
        sheet['B3'] = '=SUM(B2)'
        content = io.BytesIO()
        workbook.save(content)
        job = self.request(base + '/templates', content=content.getvalue(), headers={
            'Content-Type': 'application/octet-stream', 'X-Filename': 'budget.xlsx',
        })
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        template = self.request(base + '/templates')[0]
        with self.assertRaises(HTTPError) as unplanned:
            self.request(base + '/mappings', {'solution_id': solution['id'], 'template_id': template['source']['id']})
        self.assertEqual(unplanned.exception.code, 409)
        proposal = loads(proposal_json())
        proposal['plan']['budget_items'] = [{
            'template_id': template['source']['id'], 'sheet': '资金计划', 'cell': 'B2',
            'label': '借款', 'period': 1, 'unit': '元', 'value': '10',
            'action_indices': [0], 'explanation': '第一期借款',
        }]
        self.llm.response = dumps(proposal)
        job = self.request(base + '/solve', {'version': 1, 'generation': {'template_ids': [template['source']['id']]}})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        solution = self.request(base + '/solutions')[-1]
        entry = {'sheet': '资金计划', 'cell': 'B2', 'label': '借款',
                 'source_ids': ['action:0:amount'], 'conversion': 'identity', 'explanation': '第一期操作金额'}
        self.llm.response = dumps({'entries': [entry]})
        job = self.request(base + '/mappings', {'solution_id': solution['id'], 'template_id': template['source']['id']})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'needs_input')
        mapping = self.request(base + '/mappings')[0]
        self.assertEqual(mapping['values'], ['10'])
        with self.assertRaises(HTTPError) as changed_budget:
            self.request(base + '/mappings/' + mapping['id'], {
                'expected_revision': 1, 'entries': [{**entry, 'source_ids': []}],
            })
        self.assertEqual(changed_budget.exception.code, 400)
        with self.assertRaises(HTTPError) as unconfirmed:
            self.request(base + '/artifacts', {'solution_id': solution['id'], 'kind': 'xlsx', 'mapping_id': mapping['id']})
        self.assertEqual(unconfirmed.exception.code, 409)
        mapping = self.request(base + '/mappings/' + mapping['id'], {'expected_revision': 1, 'entries': [entry]})
        self.assertTrue(mapping['confirmed'])
        with self.assertRaises(HTTPError) as conflict:
            self.request(base + '/mappings/' + mapping['id'], {'expected_revision': 1, 'entries': [entry]})
        self.assertEqual(conflict.exception.code, 409)
        for kind, mapping_id in [('docx', None), ('xlsx', mapping['id']), ('xlsx', None)]:
            job = self.request(base + '/artifacts', {'solution_id': solution['id'], 'kind': kind, 'mapping_id': mapping_id})
            self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        artifacts = self.request(base + '/artifacts')
        self.assertEqual(len(artifacts), 3)
        for artifact in artifacts:
            code, body = self.raw_request(base + f"/artifacts/{artifact['id']}/download")
            self.assertEqual(code, 200)
            self.assertIn('Content-Disposition: attachment;', self.response_headers)
            self.assertTrue(artifact['draft'])
            if artifact['kind'] == 'docx':
                document = Document(io.BytesIO(body))
                self.assertIn('第 1 年 · 第 2 季', [p.text for p in document.paragraphs])
            elif artifact['mapping_id']:
                filled = load_workbook(io.BytesIO(body))
                self.assertEqual(filled['资金计划']['B2'].value, 10)
                self.assertEqual(filled['资金计划']['B3'].value, '=SUM(B2)')
                filled.close()
        other = self.request('/api/projects', {'name': '其他项目'})['id']
        self.assertEqual(self.raw_request(f"/api/projects/{other}/artifacts/{artifacts[0]['id']}/download")[0], 404)
        self.cookie = ''
        self.assertEqual(self.raw_request(base + f"/artifacts/{artifacts[0]['id']}/download")[0], 401)

    def test_upload_read_revise_solve_and_isolation(self):
        project, result = self.setup_problem()
        base = f"/api/projects/{project}"
        snapshot = result["snapshot"]
        self.assertEqual(snapshot["version"], 1)
        editable = {"title": "修正后的题目", "facts": snapshot["reading"]["problem"]["facts"],
                    "rules": [], "issues": []}
        revised = self.request(base + "/revisions", {"expected_version": 1,
                                                      "reading": editable, "note": "复核题目名称"})
        self.assertEqual(revised["version"], 2)
        self.assertEqual(self.request(base + "/problem?version=1")["snapshot"]["reading"]["problem"]["title"],
                         "测试题目")
        with self.assertRaises(HTTPError) as conflict:
            self.request(base + "/revisions", {"expected_version": 1, "reading": editable, "note": "旧版"})
        self.assertEqual(conflict.exception.code, 409)
        self.llm.response = proposal_json()
        job = self.request(base + "/solve", {"version": 2})
        done = self.wait_job(project, job["id"])
        self.assertEqual(done["status"], "succeeded")
        self.assertFalse(done["result"]["verified"])
        solutions = self.request(base + "/solutions")
        self.assertEqual(solutions[0]["problem_version"], 2)
        self.assertEqual(solutions[0]["mode"], "llm_only")
        with self.assertRaises(HTTPError) as mc:
            self.request(base + "/solve", {"version": 2, "mode": "llm_monte_carlo"})
        self.assertEqual(mc.exception.code, 409)
        other = self.request("/api/projects", {"name": "另一个项目"})["id"]
        with self.assertRaises(HTTPError) as isolated:
            self.request(f"/api/projects/{other}/jobs/{job['id']}")
        self.assertEqual(isolated.exception.code, 404)

    def test_legacy_false_conflict_is_rechecked_without_rewriting_snapshot(self):
        project, result = self.setup_problem()
        base = f"/api/projects/{project}"
        snapshot = result["snapshot"]
        document_id = snapshot["document_ids"][0]
        parser = OOXMLDocumentParser()
        content = docx_bytes('<w:p><w:r><w:t>每季付息，到期还 本</w:t></w:r></w:p>')
        document = parser.parse(SourceDocument(document_id, 'topic.docx', '', len(content), 'docx'), content)
        self.server.workflow.repository.put('document', project, document_id, document)
        fact = {'id': 'repayment', 'category': 'finance', 'entity_id': 'bank',
                'entity_name': '长期银行融资', 'key': 'repayment_method',
                'value': '每季付息，到期还本', 'value_type': 'text', 'unit': '无', 'state': 'known',
                'evidence_ids': [document.evidence[0].id], 'conversion': 'identity', 'explanation': ''}
        snapshot['reading']['raw_response'] = dumps({'title': '测试', 'facts': [fact], 'rules': [], 'issues': []})
        snapshot['reading']['problem']['facts'] = [{**fact, 'state': 'conflicted'}]
        snapshot['reading']['issues'] = [{'code': 'validation.unsupported_value', 'severity': 'error',
                                         'target': fact['id'], 'message': '旧校验误报',
                                         'evidence_ids': fact['evidence_ids']}]
        saved = self.server.workflow.repository.save_snapshot(project, snapshot, 1)
        current = self.request(base + '/problem')['snapshot']['reading']
        self.assertEqual(current['problem']['facts'][0]['state'], 'known')
        self.assertFalse(any(i['target'] == fact['id'] for i in current['issues']))
        self.assertEqual(self.server.workflow.repository.snapshot(project, 2), saved)
        snapshot['reading']['problem']['facts'][0]['explanation'] = '用户修改后仍有争议'
        self.server.workflow.repository.save_snapshot(project, snapshot, 2)
        current = self.request(base + '/problem')['snapshot']['reading']
        self.assertEqual(current['problem']['facts'][0]['state'], 'conflicted')
        self.assertTrue(any(i['target'] == fact['id'] and i['severity'] == 'warning' for i in current['issues']))

    def test_parse_retry_uses_stored_original(self):
        project = self.request("/api/projects", {"name": "重试"})["id"]
        base = f"/api/projects/{project}"
        job = self.request(base + "/documents", content=docx_bytes(), headers={
            "Content-Type": "application/octet-stream", "X-Filename": "topic.docx",
        })
        done = self.wait_job(project, job["id"])
        document_id = done["result"]["document_id"]
        retry = self.request(base + f"/documents/{document_id}/retry", {})
        self.assertEqual(self.wait_job(project, retry["id"])["status"], "succeeded")
        self.assertEqual(len(self.request(base + "/documents")), 1)

    def test_manual_inputs_annotations_sources_and_revision_isolation(self):
        project = self.request('/api/projects', {'name': '手动题目'})['id']
        base = f'/api/projects/{project}'
        inputs = self.request(base + '/inputs', {'revision': 0, 'problem_text': '初始现金100000元',
            'supplement_text': '贷款上限20000元', 'annotations': []})
        self.assertEqual(inputs['revision'], 1)
        def output(value='100000'):
            return dumps({'title': '手动题目', 'facts': [{'id': 'cash', 'category': 'initial',
                'entity_id': 'company', 'entity_name': '公司', 'key': 'initial_cash', 'value': value,
                'value_type': 'number', 'unit': '元', 'state': 'known',
                'evidence_ids': [manual_id], 'conversion': 'identity', 'explanation': ''}],
                'rules': [], 'issues': []})
        manual_id = text_source('手动题目', inputs['problem_text'], 'manual_').evidence[0].id
        self.llm.response = output()
        job = self.request(base + '/read', {'document_ids': [], 'input_revision': 1})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        snapshot = self.request(base + '/problem')['snapshot']
        self.assertEqual(snapshot['reading']['problem']['facts'][0]['state'], 'known')
        self.assertEqual(len(snapshot['additional_sources']), 2)
        reading = {'title': '手动题目', 'facts': snapshot['reading']['problem']['facts'], 'rules': [], 'issues': []}
        revised = self.request(base + '/revisions', {'expected_version': 1, 'reading': reading})
        self.assertEqual(revised['additional_sources'], snapshot['additional_sources'])
        job = self.request(base + '/documents', content=Path(__file__).with_name('fixtures').joinpath('orders.xls').read_bytes(),
                           headers={'Content-Type': 'application/octet-stream', 'X-Filename': 'orders.xls'})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        document_id = self.request(base + '/documents')[0]['source']['id']
        inputs['annotations'] = [{'document_id': document_id, 'text': '订单数量不得超过详单。'}]
        saved = self.request(base + '/inputs', inputs)
        with self.assertRaises(HTTPError) as stale:
            self.request(base + '/inputs', inputs)
        self.assertEqual(stale.exception.code, 409)
        with self.assertRaises(HTTPError) as changed:
            self.request(base + '/solve', {'version': revised['version']})
        self.assertEqual(changed.exception.code, 409)
        self.assertEqual(self.request(base + '/problem?version=1')['snapshot']['inputs']['revision'], 1)
        self.llm.response = output()
        job = self.request(base + '/read', {'document_ids': [document_id], 'input_revision': saved['revision']})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        self.assertIn('订单数量不得超过详单', self.llm.calls[-1].user)
        self.assertIn('贷款上限20000', self.llm.calls[-1].user)

    def test_historical_version_can_have_large_model_response(self):
        project, result = self.setup_problem()
        repository = self.server.workflow.repository
        snapshot = result["snapshot"]
        snapshot["reading"]["raw_response"] = " " * 60000
        repository.save_snapshot(project, snapshot, 1)
        self.llm.response = proposal_json()
        job = self.request(f"/api/projects/{project}/solve", {"version": 2})
        self.assertEqual(self.wait_job(project, job["id"])["status"], "succeeded")

    def test_config_does_not_expose_key_and_static(self):
        status = self.request("/api/status")
        self.assertEqual(status["solver_modes"], ["llm_only", "llm_monte_carlo"])
        self.assertNotIn("llm_api_key", status)
        code, raw = self.raw_request("/")
        self.assertEqual(code, 200)
        html = raw.decode()
        self.assertIn('id="app"', html)
        self.assertIn('type="module"', html)
        for path, content_type in [("/app.js", "text/javascript"), ("/style.css", "text/css")]:
            code, body = self.raw_request(path)
            self.assertEqual(code, 200)
            self.assertIn(content_type, self.response_headers)
            self.assertGreater(len(body), 100)

    def test_invalid_upload_and_cross_origin(self):
        project = self.request("/api/projects", {"name": "测试"})["id"]
        with self.assertRaises(HTTPError):
            self.request(f"/api/projects/{project}/documents", content=b"x", headers={
                "Content-Type": "application/octet-stream", "X-Filename": "../secret.docx",
            })
        with self.assertRaises(HTTPError) as cross:
            self.request("/api/projects", {"name": "跨站"}, headers={
                "Content-Type": "application/json", "Origin": "https://unrelated.example",
            })
        self.assertEqual(cross.exception.code, 403)

    def test_restart_persistence(self):
        project = self.request("/api/projects", {"name": "持久化"})["id"]
        self.server.workflow.repository.put("job", project, "c" * 32,
                                            {"id": "c" * 32, "project_id": project, "status": "running"})
        repository = SQLiteRepository(Path(self.temp.name) / "erp.sqlite3")
        self.assertEqual(repository.get("job", project, "c" * 32)["status"], "failed")
        self.assertEqual(repository.get("project", project, project)["name"], "持久化")

    def test_login_required_remember_logout_and_restart(self):
        saved_cookie = self.cookie
        self.cookie = ""
        code, _ = self.raw_request("/api/projects")
        self.assertEqual(code, 401)
        with self.assertRaises(HTTPError):
            self.request("/api/auth/login", {"key": "wrong"})
        self.request("/api/auth/login", {"key": "test-access"})
        self.assertIn("Max-Age=2592000", self.response_headers)
        self.assertIn("HttpOnly", self.response_headers)
        self.assertIn("SameSite=Strict", self.response_headers)
        self.server.access = create_access(self.server.settings)
        self.assertTrue(self.request("/api/auth/session")["authenticated"])
        self.request("/api/auth/logout", {})
        self.assertFalse(self.request("/api/auth/session")["authenticated"])
        self.cookie = saved_cookie
        self.assertTrue(self.request("/api/auth/session")["authenticated"])

    def test_session_cookie_without_remember_and_key_rotation(self):
        self.request("/api/auth/login", {"key": "test-access", "remember": False})
        self.assertNotIn("Max-Age", self.response_headers)
        self.server.access = create_access(replace(self.server.settings, access_key="rotated"))
        self.assertFalse(self.request("/api/auth/session")["authenticated"])

    def test_model_settings_persist_and_never_return_secret(self):
        data = self.request("/api/model-settings", {"base_url": "https://example.com/v1",
                                                    "model": "custom-model", "api_key": "secret-123"})
        self.assertTrue(data["api_key_set"])
        self.assertTrue(data["configured"])
        self.assertNotIn("secret-123", dumps(self.request("/api/model-settings")))
        self.request("/api/model-settings", {"model": "another", "api_key": ""})
        from erp_agent.adapters.model_settings import ModelSettingsService
        settings = ModelSettingsService(self.server.settings)
        self.assertEqual(settings.public()["model"], "another")
        self.assertEqual(settings.settings.llm_api_key, "secret-123")
        self.request("/api/model-settings", {"clear_api_key": True})
        self.assertFalse(self.request("/api/model-settings")["api_key_set"])

    def test_provider_setting_is_visible_and_switches_runtime_status(self):
        metadata = self.request("/api/model-settings")
        self.assertEqual(len(metadata["providers"]), 5)
        changed = self.request("/api/model-settings", {"provider": "ollama", "model": "local-model"})
        self.assertEqual(changed["base_url"], "http://localhost:11434")
        self.assertFalse(changed["api_key_set"])
        self.assertTrue(self.request("/api/status")["llm_configured"])
        self.request("/api/model-settings", {"provider": "gemini", "model": "gemini-model"})
        self.assertFalse(self.request("/api/status")["llm_configured"])

    def test_preset_connection_endpoint_and_authentication(self):
        self.request("/api/model-settings",{"model":"model","api_key":"secret"})
        with patch("erp_agent.adapters.model_settings.LangChainLLMClient.generate",
                   return_value=LLMResponse('{"status":"ok"}',"model")):
            result=self.request("/api/model-presets/default/test",{"expected_revision":2})
        self.assertTrue(result["succeeded"])
        self.cookie=""
        self.assertEqual(self.raw_request("/api/model-presets/default/test",method="POST")[0],401)

    def test_preset_endpoints_create_edit_switch_delete_and_authentication(self):
        catalog = self.request("/api/model-presets")
        self.assertEqual(catalog["active_id"], "default")
        preset = self.request("/api/model-presets", {
            "name": "本地预设", "provider": "ollama", "model": "local-model",
        })
        self.assertEqual(self.request("/api/model-presets")["active_id"], "default")
        preset = self.request(f"/api/model-presets/{preset['id']}", {
            "expected_revision": 1, "name": "生产备用",
        })
        self.request(f"/api/model-presets/{preset['id']}/activate", {"expected_revision": 2})
        self.assertTrue(self.request("/api/status")["llm_configured"])
        path = f"/api/model-presets/{preset['id']}?revision=2"
        self.assertEqual(self.raw_request(path, method="DELETE")[0], 409)
        self.request("/api/model-presets/default", {"expected_revision": 1,
                                                     "model": "primary", "api_key": "secret"})
        self.request("/api/model-presets/default/activate", {"expected_revision": 2})
        self.assertEqual(self.raw_request(f"/api/model-presets/{preset['id']}?revision=2",
                                         method="DELETE")[0], 200)
        self.cookie = ""
        self.assertEqual(self.raw_request("/api/model-presets")[0], 401)
        self.assertEqual(self.raw_request("/api/model-presets/default/activate", method="POST")[0], 401)

    def test_queued_read_and_solve_retain_submission_preset(self):
        project, _ = self.setup_problem()
        workflow = self.server.workflow
        runner = workflow.jobs
        runner.executor.shutdown(wait=True)
        runner.executor = ThreadPoolExecutor(max_workers=1)
        entered, release = threading.Event(), threading.Event()

        def block(control):
            entered.set()
            release.wait(3)
            return {}

        held = runner.submit(project, "test", block)
        self.assertTrue(entered.wait(2))
        self.llm.response = proposal_json()
        read = self.request(f"/api/projects/{project}/read", {})
        solve = self.request(f"/api/projects/{project}/solve", {"version": 1})
        original_response = self.llm.response
        original_generate = self.llm.generate

        def generate(request):
            if request.schema_name == "erp_reading":
                snapshot = workflow.repository.snapshot(project, 1)
                self.llm.response = dumps({"title": "题目", "facts": snapshot["reading"]["problem"]["facts"],
                                           "rules": [], "issues": []})
            else:
                self.llm.response = original_response
            return original_generate(request)

        self.llm.generate = generate
        try:
            preset = self.request("/api/model-presets", {
                "name": "切换后模型", "provider": "ollama", "model": "replacement",
            })
            self.request(f"/api/model-presets/{preset['id']}/activate", {"expected_revision": 1})
            release.set()
            self.wait_job(project, held["id"])
            self.assertEqual(self.wait_job(project, read["id"])["status"], "succeeded")
            self.assertEqual(self.wait_job(project, solve["id"])["status"], "succeeded")
            self.assertEqual(read["metadata"]["model_preset"]["preset_id"], "default")
            self.assertEqual(solve["metadata"]["model_preset"]["preset_revision"], 1)
            self.assertEqual(self.request(f"/api/projects/{project}/solutions")[0]["model"], "test-model")
        finally:
            release.set()

    def test_delete_project_cleans_every_record_and_blocks_late_write(self):
        project, result = self.setup_problem()
        repository = self.server.workflow.repository
        repository.put("solution", project, "s", {"id": "s"})
        code, _ = self.raw_request(f"/api/projects/{project}", method="DELETE")
        self.assertEqual(code, 200)
        self.assertEqual(repository.list("original", project), [])
        self.assertEqual(repository.list("document", project), [])
        self.assertEqual(repository.list("solution", project), [])
        self.assertEqual(repository.versions(project), [])
        from erp_agent.domain.errors import ERPError
        with self.assertRaises(ERPError):
            repository.put("document", project, "late", {"id": "late"})

    def test_selected_documents_must_belong_to_project(self):
        project, _ = self.setup_problem()
        with self.assertRaises(HTTPError) as error:
            self.request(f"/api/projects/{project}/read", {"document_ids": ["not-owned"]})
        self.assertEqual(error.exception.code, 400)

    def test_delete_running_project_does_not_restore_data(self):
        project = self.request("/api/projects", {"name": "运行中删除"})["id"]
        entered, release = threading.Event(), threading.Event()
        repository = self.server.workflow.repository

        def operation(control):
            entered.set()
            release.wait(3)
            control.checkpoint("saving_document")
            repository.put("document", project, "late", {"id": "late"})
            return {"ok": True}

        runner = self.server.workflow.jobs
        job = runner.submit(project, "parse", operation)
        try:
            self.assertTrue(entered.wait(2))
            code, _ = self.raw_request(f"/api/projects/{project}", method="DELETE")
            self.assertEqual(code, 200)
            release.set()
            runner.futures[job["id"]].result(timeout=3)
            self.assertEqual(repository.list("job", project), [])
            self.assertEqual(repository.list("document", project), [])
        finally:
            release.set()

    def test_selected_read_uses_only_selected_document_and_blocks_duplicate(self):
        project, _ = self.setup_problem()
        base = f"/api/projects/{project}"
        job = self.request(base + "/documents", content=docx_bytes(), headers={
            "Content-Type": "application/octet-stream", "X-Filename": "second.docx",
        })
        self.wait_job(project, job["id"])
        document_id = self.request(base + "/documents")[0]["source"]["id"]
        original = self.server.workflow.reader
        entered, release = threading.Event(), threading.Event()

        class DelayedReader:
            llm = original.llm

            def read(self, documents, control, client=None):
                entered.set()
                release.wait(3)
                return original.read(documents, control, client)

        self.server.workflow.reader = DelayedReader()
        try:
            job = self.request(base + "/read", {"document_ids": [document_id]})
            self.assertTrue(entered.wait(2))
            with self.assertRaises(HTTPError) as duplicate:
                self.request(base + "/read", {"document_ids": [document_id]})
            self.assertEqual(duplicate.exception.code, 409)
            release.set()
            self.assertEqual(self.wait_job(project, job["id"])["status"], "succeeded")
            self.assertEqual(self.request(base + "/problem")["snapshot"]["document_ids"],
                             [document_id])
        finally:
            release.set()

    def test_sensitive_endpoints_need_login_and_secure_cookie_is_optional(self):
        project = self.request("/api/projects", {"name": "受保护"})["id"]
        self.cookie = ""
        for path, method in [("/api/model-settings", "GET"),
                             ("/api/model-settings", "POST"),
                             (f"/api/projects/{project}", "DELETE")]:
            self.assertEqual(self.raw_request(path, method=method)[0], 401)
        self.server.settings = replace(self.server.settings, cookie_secure=True)
        self.request("/api/auth/login", {"key": "test-access"})
        self.assertIn("; Secure", self.response_headers)

    def test_document_delete_method_is_rejected(self):
        project = self.request("/api/projects", {"name": "方法检查"})["id"]
        self.assertEqual(self.raw_request(f"/api/projects/{project}/documents", method="DELETE")[0],
                         405)

    def test_reading_errors_block_solving_before_model_call(self):
        project, result = self.setup_problem()
        reading = result["snapshot"]["reading"]
        facts = reading["problem"]["facts"]
        facts[0]["value"] = "72000"
        revision = self.request(f"/api/projects/{project}/revisions", {
            "expected_version": 1, "note": "测试来源不支持的金额",
            "reading": {"title": "错误值", "facts": facts, "rules": [], "issues": []},
        })
        self.assertTrue(any(i["severity"] == "error" for i in revision["reading"]["issues"]))
        calls = len(self.llm.calls)
        with self.assertRaises(HTTPError) as invalid:
            self.request(f"/api/projects/{project}/solve", {"version": revision["version"]})
        self.assertEqual(invalid.exception.code, 409)
        self.assertEqual(len(self.llm.calls), calls)

    def test_confirmed_rate_basis_and_answer_reach_solver_without_changing_rate(self):
        project, result = self.setup_problem()
        base = f'/api/projects/{project}'
        reading = result['snapshot']['reading']
        fact = {'id': 'rate', 'category': 'finance', 'entity_id': 'direct',
                'entity_name': '直接融资', 'key': 'interest_rate', 'value': '0.05',
                'value_type': 'number', 'unit': '比例', 'state': 'assumed',
                'evidence_ids': [], 'conversion': 'percent_to_ratio',
                'explanation': '利率5%转为0.05；用户答复：对应贷款期限的总利率。'}
        basis = {**fact, 'id': 'basis', 'key': 'interest_rate_basis', 'value': 'term_total',
                 'value_type': 'text', 'unit': '', 'conversion': 'identity',
                 'explanation': '用户确认：对应完整期限的总利率。'}
        issue = {'code': 'review.resolved.rate_period', 'target': 'rate', 'severity': 'info',
                 'message': '利率覆盖哪个期限？；用户答复：对应完整贷款期限。', 'evidence_ids': []}
        revised = self.request(base + '/revisions', {'expected_version': 1,
            'reading': {'title': '计息口径', 'facts': reading['problem']['facts'] + [fact, basis],
                        'rules': [], 'issues': [issue]}})
        self.assertEqual(self.request(base + '/problem')['snapshot']['reading']['problem']['facts'][-2:],
                         [fact, basis])
        self.llm.response = proposal_json()
        job = self.request(base + '/solve', {'version': revised['version']})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
        material = loads(self.llm.calls[-1].user)['snapshot']
        self.assertEqual(material['problem']['facts'][-2:], [fact, basis])
        self.assertIn(issue, material['issues'])

    def test_missing_horizon_uses_defaults_and_generates_without_confirmation(self):
        project, result = self.setup_problem()
        base = f'/api/projects/{project}'
        facts = [f for f in result['snapshot']['reading']['problem']['facts']
                 if f['key'] == 'initial_cash']
        revised = self.request(base + '/revisions', {
            'expected_version': 1, 'note': '',
            'reading': {'title': '未提供周期', 'facts': facts, 'rules': [], 'issues': [
                {'code': 'missing_periods', 'target': 'initial.periods_per_year',
                 'severity': 'error', 'message': '经营年限未提供，需补充经营周期。', 'evidence_ids': []},
            ]},
        })
        self.assertEqual(revised['revision_note'], '默认')
        problem = revised['reading']['problem']
        self.assertEqual((problem['years'], problem['periods_per_year']), (4, 4))
        self.assertFalse(any(i['severity'] == 'error' for i in revised['reading']['issues']))
        self.assertEqual(self.request(base + '/problem')['snapshot']['reading']['problem'], problem)
        proposal = loads(proposal_json())
        proposal['plan']['actions'] = [
            {**proposal['plan']['actions'][0], 'period': period} for period in range(1, 17)]
        proposal['parameters'] = []
        self.llm.response = dumps(proposal)
        job = self.request(base + '/solve', {'version': revised['version']})
        self.assertEqual(self.wait_job(project, job['id'])['status'], 'succeeded')
