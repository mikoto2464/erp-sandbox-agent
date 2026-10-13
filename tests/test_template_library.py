import io
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook

from erp_agent.adapters.budget_templates import SpreadsheetBudgetParser
from erp_agent.adapters.persistence import SQLiteRepository
from erp_agent.api.server import create_workflow
from erp_agent.application.contracts import dumps, loads, obj, to_json
from erp_agent.application.ports import LLMResponse
from erp_agent.application.template_reader import TemplateReader, template_material
from erp_agent.config import Settings
from erp_agent.domain.errors import ERPError
from erp_agent.domain.models import SourceDocument
from tests.helpers import Control, FakeLLM, snapshot


def workbook_bytes(rows=24):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预算"
    sheet.append(["项目", "第一季"])
    for index in range(2, rows + 1):
        sheet.cell(index, 1, "原料支出")
        sheet.cell(index, 2).number_format = "0.00"
    sheet.cell(rows + 1, 2, f"=SUM(B2:B{rows})")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def output(target="B2"):
    return dumps(
        {
            "purpose": "现金预算",
            "periods": "第一季",
            "input_groups": [
                {"label": "原料支出", "range": target, "unit": "元", "description": "现金支出输入"}
            ],
            "issues": [],
        }
    )


class TemplateReaderTests(unittest.TestCase):
    def test_gateway_failure_splits_then_restart_reuses_completed_segments(self):
        content = workbook_bytes()
        template = SpreadsheetBudgetParser().parse(
            SourceDocument("template", "budget.xlsx", "hash", len(content), "xlsx"), content
        )
        calls, plans = [], []

        class Client:
            configured = True
            failing = True

            def plan_reading(self, requests, control, concurrency):
                plans.append([r.task.id for r in requests])

            def generate(self, request):
                data = loads(request.user)
                calls.append(data)
                if "row_window" not in data or (data["row_window"][0] > 1 and self.failing):
                    raise ERPError("llm_gateway_timeout", "HTTP 504", 504)
                first = data["row_window"][0]
                return LLMResponse(output(f"B{max(2, first)}"), "test")

        with tempfile.TemporaryDirectory() as directory:
            repository = SQLiteRepository(Path(directory) / "erp.sqlite3")
            reader, client = TemplateReader(repository), Client()
            with self.assertRaises(ERPError):
                reader.read(template, client, Control())
            first_segment_calls = sum(d.get("row_window") == [1, 13] for d in calls)
            client.failing = False
            reading = reader.read(template, client, Control())
            self.assertEqual(
                sum(d.get("row_window") == [1, 13] for d in calls), first_segment_calls
            )
            before = len(calls)
            again = reader.read(template, client, Control())
            self.assertEqual(len(calls), before)
            self.assertEqual(reading, again)
            self.assertEqual(len(plans), 2)
            material = template_material(template, template.sheets[0])
            self.assertNotIn("blank_ranges", material["sheet"])
            self.assertFalse(any(c[3] for c in material["sheet"]["cells"]))
            self.assertNotIn("current_problem_context", calls[0])

    def test_library_import_is_model_free_and_survives_project_and_library_deletion(self):
        with tempfile.TemporaryDirectory() as directory:
            client = FakeLLM(output())
            workflow = create_workflow(Settings(data_dir=Path(directory)), client)
            try:
                library = workflow.template_library
                job = library.upload("budget.xlsx", workbook_bytes())
                workflow.jobs.futures[job["id"]].result(3)
                tid = job["metadata"]["template_id"]
                job = library.read(tid)
                workflow.jobs.futures[job["id"]].result(3)
                library.review(tid, {"reviewed": True, "explanation": ""})
                calls = len(client.calls)
                p1 = workflow.create_project("一")["id"]
                p2 = workflow.create_project("二")["id"]
                problem = replace(snapshot(), project_id=p1, document_ids=())
                workflow.repository.save_snapshot(p1, obj(to_json(problem)), 0)
                imported1 = workflow.import_library_template(p1, {"template_id": tid})
                imported2 = workflow.import_library_template(p2, {"template_id": tid})
                self.assertEqual(len(client.calls), calls)
                self.assertEqual(imported1["version"], 2)
                self.assertIsNone(imported2["version"])
                self.assertTrue(
                    workflow.repository.snapshot(p1)["template_readings"][0]["reviewed"]
                )
                self.assertNotEqual(imported1["template_id"], imported2["template_id"])
                workflow.delete_project(p1)
                self.assertEqual(len(library.items()), 1)
                library.delete(tid)
                self.assertEqual(library.items(), [])
                preserved = workflow.repository.get("template", p2, imported2["template_id"])
                self.assertTrue(any(c["formula"] for s in preserved["sheets"] for c in s["cells"]))
                self.assertTrue(
                    workflow.repository.get(
                        "imported_template_reading", p2, imported2["template_id"]
                    )["reviewed"]
                )
            finally:
                workflow.jobs.close()
                workflow.memory.close()
